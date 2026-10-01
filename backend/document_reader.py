"""Bounded document readers and local OCR. Never decode binary files as plain text."""
from __future__ import annotations

import base64
import os
import io
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import zipfile

from backend.config import MAX_RESUME_BYTES


class DocumentReadError(ValueError):
    """A safe, actionable error which may be displayed to the recruiter."""


IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif")
MAX_PAGES = 30
MAX_EXPANDED_BYTES = 40 * 1024 * 1024


def _text(data: bytes) -> str:
    for encoding in (("utf-16",) if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "cp1252")):
        try:
            text = data.decode(encoding).replace("\x00", "")
            if sum(ord(c) < 32 and c not in "\n\r\t" for c in text) > max(1, len(text) // 100):
                break
            return text
        except UnicodeError:
            continue
    raise DocumentReadError("This document contains unreadable text. Export it as PDF or DOCX and retry.")


def _check_archive(data: bytes):
    archive = zipfile.ZipFile(io.BytesIO(data))
    if sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED_BYTES:
        archive.close()
        raise DocumentReadError("The expanded document is too large. Upload a smaller document.")
    return archive


def read_document(filename: str, data: bytes) -> str:
    if not data:
        raise DocumentReadError("The file is empty. Upload a readable copy.")
    if len(data) > MAX_RESUME_BYTES:
        raise DocumentReadError(f"File exceeds the {MAX_RESUME_BYTES // (1024 * 1024)} MB document limit.")
    suffix = Path(filename).suffix.lower()
    try:
        if data.startswith(b"%PDF"):
            return _read_pdf(data)
        if data.startswith(b"PK"):
            with _check_archive(data) as archive:
                names = archive.namelist()
                if "word/document.xml" in names:
                    from docx import Document
                    doc = Document(io.BytesIO(data))
                    lines = [p.text for p in doc.paragraphs]
                    lines += [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
                    for section in doc.sections:
                        lines += [p.text for p in section.header.paragraphs + section.footer.paragraphs]
                    return "\n".join(lines)
                if "content.xml" in names:
                    from defusedxml.ElementTree import fromstring
                    root = fromstring(archive.read("content.xml"))
                    return "\n".join("".join(el.itertext()) for el in root.iter()
                                     if el.tag.endswith(("}p", "}h")))
            raise DocumentReadError("This archive is not a supported Word or OpenDocument file.")
        if data.startswith(b"{\\rtf"):
            from striprtf.striprtf import rtf_to_text
            return rtf_to_text(data.decode("cp1252"))
        if data.startswith(b"\xd0\xcf\x11\xe0"):
            command = shutil.which("antiword")
            mac_reader = shutil.which("textutil") if sys.platform == "darwin" else None
            if not command and not mac_reader:
                raise DocumentReadError("Legacy Word reader is unavailable. Ask your administrator to install antiword, or save as DOCX.")
            with tempfile.TemporaryDirectory(prefix="resume-doc-") as directory:
                path = Path(directory) / "document.doc"
                path.write_bytes(data)
                args = [command, "-m", "UTF-8.txt", str(path)] if command else [mac_reader, "-convert", "txt", "-stdout", str(path)]
                result = subprocess.run(args, capture_output=True, timeout=20)
            if result.returncode:
                raise DocumentReadError("This Word file could not be read. Save an unlocked copy as DOCX or PDF.")
            return result.stdout.decode("utf-8", errors="replace")
        if suffix in IMAGE_EXTENSIONS or data.startswith((b"\x89PNG", b"\xff\xd8\xff", b"II*\x00", b"MM\x00*", b"GIF8", b"BM")):
            # Mail logos, warning badges, and signature icons are frequent
            # attachments. Ignore tiny images before invoking local or cloud OCR.
            from PIL import Image
            with Image.open(io.BytesIO(data)) as image:
                if image.width < 500 or image.height < 600:
                    return ""
            return _ocr(data, "image")
        # Preserve text fixtures/exported text with a .pdf/.docx suffix, but never
        # treat known binary formats as text after their parser fails.
        if suffix in (".txt", ".md", ".text", ".pdf", ".docx", ".doc", ".rtf", ".odt"):
            return _text(data)
        raise DocumentReadError("Unsupported file format. Use PDF, DOCX, DOC, ODT, RTF, text, or a résumé image.")
    except DocumentReadError:
        raise
    except subprocess.TimeoutExpired:
        raise DocumentReadError("Document processing timed out. Split the document or upload a clearer copy.") from None
    except Exception:
        raise DocumentReadError("The document is damaged, locked, or unreadable. Upload an unlocked PDF or DOCX copy.") from None


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentReadError("This PDF is password protected. Upload an unlocked copy.")
    if len(reader.pages) > MAX_PAGES:
        raise DocumentReadError(f"Document exceeds {MAX_PAGES} pages. Split it into individual résumés.")
    pages = []
    scan_pages = []
    for index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages.append(text)
        # Inspect each page independently: mixed digital/scanned PDFs need OCR too.
        if len("".join(c for c in text if c.isalnum())) < 40:
            scan_pages.append(index)
    if scan_pages:
        import json
        recognized = json.loads(_ocr(data, "pdf", scan_pages))
        for index in scan_pages:
            pages[index] = recognized.get(str(index), "") or pages[index]
    return "\n".join(pages)


def _ocr(data: bytes, kind: str, pages: list[int] | None = None) -> str:
    command = os.getenv("TESSERACT_CMD", "tesseract")
    api_key = os.getenv("GOOGLE_VISION_API_KEY", "").strip()
    if not shutil.which(command):
        if api_key:
            return _google_vision_ocr(data, kind, pages)
        raise DocumentReadError(
            "OCR is unavailable. Install Tesseract on the server or set the optional "
            "GOOGLE_VISION_API_KEY server secret, then retry."
        )
    # Isolate PDFium and image decoding from concurrent request threads. The
    # parent imposes a whole-document deadline, in addition to per-page limits.
    try:
        result = subprocess.run(
            [sys.executable, "-m", "backend.ocr_worker", kind, ",".join(map(str, pages or []))],
            input=data, capture_output=True, timeout=60,
        )
        if result.returncode:
            raise DocumentReadError(
                "Local OCR could not read this file. Check the server's OCR language "
                "configuration or upload a clearer copy."
            )
        return result.stdout.decode("utf-8")
    except (DocumentReadError, subprocess.TimeoutExpired):
        if api_key:
            return _google_vision_ocr(data, kind, pages)
        raise


def _google_vision_ocr(data: bytes, kind: str, pages: list[int] | None = None) -> str:
    """Opt-in OCR fallback; only invoked when GOOGLE_VISION_API_KEY is configured."""
    import httpx

    images: list[bytes]
    if kind == "pdf":
        import pypdfium2 as pdfium
        images = []
        with pdfium.PdfDocument(data) as document:
            for index in pages or []:
                page = document[index]
                width, height = page.get_size()
                scale = min(2.5, (20_000_000 / max(1, width * height)) ** 0.5)
                bitmap = page.render(scale=scale)
                image = bitmap.to_pil()
                try:
                    output = io.BytesIO()
                    image.save(output, format="JPEG", quality=85)
                    images.append(output.getvalue())
                finally:
                    image.close()
                    bitmap.close()
                    page.close()
    else:
        images = [data]

    if not images:
        return ""

    extracted: list[str] = []
    endpoint = "https://vision.googleapis.com/v1/images:annotate"
    api_key = os.environ["GOOGLE_VISION_API_KEY"].strip()
    try:
        # Vision accepts a batch of images in one annotate call. Keep batches
        # bounded and preserve page order in the returned text.
        with httpx.Client(timeout=httpx.Timeout(25, connect=8)) as client:
            for offset in range(0, len(images), 16):
                requests = [
                    {
                        "image": {"content": base64.b64encode(image).decode("ascii")},
                        "features": [{"type": "DOCUMENT_TEXT_DETECTION"}],
                    }
                    for image in images[offset : offset + 16]
                ]
                response = client.post(
                    endpoint,
                    params={"key": api_key},
                    json={"requests": requests},
                )
                if response.is_error:
                    # Avoid returning request URLs or provider payloads that may
                    # contain credential details to the UI or application logs.
                    raise DocumentReadError(
                        "Google Vision OCR failed. Check the server API key, enabled "
                        "Vision API, and project quota."
                    )
                for result in response.json().get("responses", []):
                    if result.get("error"):
                        extracted.append("")
                        continue
                    extracted.append(
                        result.get("fullTextAnnotation", {}).get("text", "")
                        or "\n".join(
                            item.get("description", "")
                            for item in result.get("textAnnotations", [])[:1]
                        )
                    )
    except DocumentReadError:
        raise
    except httpx.TimeoutException:
        raise DocumentReadError("Google Vision OCR timed out. Retry the file or upload a smaller image.") from None
    except httpx.HTTPError:
        raise DocumentReadError("Could not connect to Google Vision OCR. Retry the file shortly.") from None
    except (ValueError, KeyError):
        raise DocumentReadError("Google Vision returned an unreadable response. Retry the file.") from None

    if kind == "pdf":
        import json
        return json.dumps({str(index): text for index, text in zip(pages or [], extracted)})
    return "\n".join(extracted)
